from pathlib import Path
import re

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


BASE = Path("/Users/chenminggang/Documents/trae_projects/iot-cross-domain/docs/db_design")
MD_PATH = BASE / "database_design.md"
IMG_PATH = BASE / "er_model.png"
OUT_PATH = BASE / "database_design_for_thesis.docx"


def set_run_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def set_paragraph_font(paragraph, name="宋体", size=12, bold=False):
    if not paragraph.runs:
        paragraph.add_run("")
    for run in paragraph.runs:
        set_run_font(run, name=name, size=size, bold=bold)


def heading_font(level):
    if level == 1:
        return ("黑体", 16, True)
    if level == 2:
        return ("黑体", 14, True)
    return ("黑体", 12, True)


def parse_md_table(lines, start):
    rows = []
    i = start
    while i < len(lines):
        line = lines[i].rstrip()
        if "|" not in line or not line.strip().startswith("|"):
            break
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def clean_text(text: str) -> str:
    # Remove markdown inline code/bold marks for thesis-friendly text.
    t = text.replace("`", "")
    t = t.replace("**", "")
    return t


def normalize_rows(rows):
    if not rows:
        return []
    sep = re.compile(r"^:?-{3,}:?$")
    out = []
    for idx, r in enumerate(rows):
        if idx == 1 and all(sep.match(c.replace(" ", "")) for c in r):
            continue
        out.append([clean_text(c) for c in r])
    width = max(len(r) for r in out)
    return [r + [""] * (width - len(r)) for r in out]


def add_table(doc, rows):
    rows = normalize_rows(rows)
    if len(rows) < 2:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True

    hdr = table.rows[0].cells
    for i, cell in enumerate(hdr):
        cell.text = rows[0][i]
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_font(p, "宋体", 10.5, True)

    for r in rows[1:]:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_paragraph_font(p, "宋体", 10.5, False)
    doc.add_paragraph("")


def main():
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    i = 0
    inserted_er_image = False

    while i < len(lines):
        line = lines[i].rstrip()
        text = line.strip()
        if not text:
            i += 1
            continue

        if text.startswith("#### "):
            p = doc.add_paragraph(clean_text(text[5:]))
            f, s, b = heading_font(3)
            set_paragraph_font(p, f, s, b)
            i += 1
            continue

        if text.startswith("### "):
            p = doc.add_paragraph(clean_text(text[4:]))
            f, s, b = heading_font(2)
            set_paragraph_font(p, f, s, b)
            if "E-R 模型" in text and IMG_PATH.exists() and not inserted_er_image:
                cap = doc.add_paragraph("图4.5 全局 E-R 图")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_paragraph_font(cap, "宋体", 10.5, False)
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.add_run().add_picture(str(IMG_PATH), width=Cm(15.5))
                inserted_er_image = True
            i += 1
            continue

        if text.startswith("## "):
            p = doc.add_paragraph(clean_text(text[3:]))
            f, s, b = heading_font(1)
            set_paragraph_font(p, f, s, b)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue

        if text.startswith("|") and "|" in text:
            rows, nxt = parse_md_table(lines, i)
            add_table(doc, rows)
            i = nxt
            continue

        if text.startswith("- "):
            p = doc.add_paragraph("• " + clean_text(text[2:]))
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Pt(0)
            set_paragraph_font(p, "宋体", 12, False)
            i += 1
            continue

        if text.startswith("---"):
            i += 1
            continue

        p = doc.add_paragraph(clean_text(text))
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
        set_paragraph_font(p, "宋体", 12, False)
        i += 1

    doc.save(OUT_PATH)
    print(f"Generated: {OUT_PATH}")


if __name__ == "__main__":
    main()
